#!/usr/bin/env python3
"""資料（pptx）の文字を、スライドごと・図形ごとに一覧で出す。

直したい文字の「図形の番号」と「いまの文字」を調べるために使う。
その2つが分かれば、edit_text.py で差し替えられる。

    python3 docs/deck/dump_text.py static/guide/00_概要.pptx
    python3 docs/deck/dump_text.py static/guide/00_概要.pptx 3      # 3枚目だけ

必要なもの: pip install python-pptx
"""

import sys

from pptx import Presentation


def main(path, only=None):
    pres = Presentation(path)
    for no, slide in enumerate(pres.slides, 1):
        if only and no != only:
            continue
        print(f'===== スライド {no}')
        for index, shape in enumerate(slide.shapes):
            if shape.has_text_frame and shape.text_frame.text.strip():
                print(f'  [{index}] {shape.text_frame.text!r}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    main(sys.argv[1], int(sys.argv[2]) if len(sys.argv) > 2 else None)
