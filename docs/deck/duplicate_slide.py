#!/usr/bin/env python3
"""資料（pptx）のスライドを複製して、好きな位置に入れる。

枚数を増やしたいときに使う。似た作りのスライドを複製し、そのあと edit_text.py で
中身を入れ替える（生成スクリプトが無いため、この手順で足す）。

    python3 docs/deck/duplicate_slide.py <pptx> <複製するスライド番号> <入れる位置>

    # 例: 9枚目を複製して、10枚目の直後（＝11枚目）に入れる
    python3 docs/deck/duplicate_slide.py static/guide/01_リポジトリの作り方と運用.pptx 9 11

複製したスライドの文字は複製元のままなので、続けて次を行う。

    python3 docs/deck/dump_text.py <pptx> <入れた位置>      # 図形の番号を調べる
    python3 docs/deck/edit_text.py docs/deck/changes/<名前>.json

アイコン（画像）も複製元のものが入る。意味が合わないときは、同じ資料の別の
スライドから画像を持ってきて差し替える（README の「アイコンを差し替える」）。

必要なもの: pip install python-pptx
"""

import copy
import sys

from pptx import Presentation

NS = '{http://schemas.openxmlformats.org/presentationml/2006/main}'


def duplicate(pres, src_index):
    """スライドを複製して末尾に足す（背景・図形・画像ごと）。"""
    src = pres.slides[src_index]
    dst = pres.slides.add_slide(src.slide_layout)

    # 中身（背景＋図形）をまるごと差し替える
    dst.element.replace(dst.element.find(f'{NS}cSld'),
                        copy.deepcopy(src.element.find(f'{NS}cSld')))

    # 画像などの参照を張り直し、複製した XML の中の rId を新しい番号に書き換える
    renamed = {}
    for rid, rel in src.part.rels.items():
        if rel.reltype.endswith('slideLayout'):
            continue
        target = rel.target_ref if rel.is_external else rel.target_part
        renamed[rid] = dst.part.rels._add_relationship(rel.reltype, target, rel.is_external)
    for elem in dst.element.iter():
        for attr, value in list(elem.attrib.items()):
            if attr.rpartition('}')[2] in ('id', 'embed', 'link') and value in renamed:
                elem.set(attr, renamed[value])


def move(pres, from_pos, to_pos):
    """スライドの並びを入れ替える（0 始まり）。"""
    lst = pres.slides._sldIdLst
    ids = list(lst)
    lst.remove(ids[from_pos])
    lst.insert(to_pos, ids[from_pos])


def main(path, source_no, position):
    pres = Presentation(path)
    duplicate(pres, source_no - 1)
    move(pres, len(pres.slides._sldIdLst) - 1, position - 1)
    pres.save(path)
    print(f'{path}: {source_no} 枚目を複製して {position} 枚目に入れました'
          f'（全 {len(pres.slides)} 枚）')


if __name__ == '__main__':
    if len(sys.argv) != 4:
        sys.exit(__doc__)
    main(sys.argv[1], int(sys.argv[2]), int(sys.argv[3]))
